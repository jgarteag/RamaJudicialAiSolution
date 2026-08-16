import io
import json
import logging
import os
import uuid
from datetime import datetime, timezone

from app.adapters.mongodb_adapter import MongoDBRadicadoRepository
from app.adapters.secrets_adapter import SecretsManagerProvider
from app.domain.services.search_service import SearchService

logger = logging.getLogger(__name__)
logger.setLevel(logging.INFO)

ENVIRONMENT = os.environ.get("ENVIRONMENT", "dev")
MONGODB_SECRET_NAME = os.environ.get("MONGODB_SECRET_NAME", "")
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt", ".text"}

_secret_provider = None
_radicado_repo = None


def _get_secret_provider() -> SecretsManagerProvider:
    global _secret_provider
    if _secret_provider is None:
        _secret_provider = SecretsManagerProvider()
    return _secret_provider


def _get_radicado_repo() -> MongoDBRadicadoRepository:
    global _radicado_repo
    if _radicado_repo is None:
        _radicado_repo = MongoDBRadicadoRepository(
            secret_provider=_get_secret_provider(),
            secret_name=MONGODB_SECRET_NAME,
        )
    return _radicado_repo


def extract_text(file_bytes: bytes, filename: str) -> str:
    ext = os.path.splitext(filename.lower())[1] if filename else ""

    if ext == ".pdf":
        from PyPDF2 import PdfReader

        reader = PdfReader(io.BytesIO(file_bytes))
        text = ""
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
        return text

    if ext in (".docx", ".doc"):
        from docx import Document

        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs)

    if ext in (".txt", ".text", ""):
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1")

    raise ValueError(f"Formato no soportado: {ext}. Usa PDF, DOCX o TXT.")


def success(body: dict) -> dict:
    return {
        "statusCode": 200,
        "headers": {"Content-Type": "application/json"},
        "body": json.dumps(body, ensure_ascii=False),
    }


def error(status: int, message: str, request_id: str = "") -> dict:
    return {
        "statusCode": status,
        "headers": {"Content-Type": "application/json"},
        "body": json.dumps({"error": message, "request_id": request_id}, ensure_ascii=False),
    }


def lambda_handler(event, context):
    request_id = str(uuid.uuid4())
    timestamp = datetime.now(timezone.utc).isoformat()

    http_method = event.get("requestContext", {}).get("http", {}).get("method", "")
    raw_path = event.get("rawPath", "")

    logger.info(
        "Request",
        extra={"request_id": request_id, "method": http_method, "path": raw_path},
    )

    if "health" in raw_path and http_method == "GET":
        return success({
            "status": "healthy",
            "service": "rama-judicial-ai",
            "environment": ENVIRONMENT,
            "formats": list(SUPPORTED_EXTENSIONS),
            "timestamp": timestamp,
        })

    if "juzgados" in raw_path and http_method == "GET":
        try:
            juzgados = _get_radicado_repo().list_juzgados()
            return success({"juzgados": juzgados, "total": len(juzgados)})
        except Exception:
            logger.error("MongoDB error", exc_info=True)
            return error(503, "No se pudo conectar a la base de datos", request_id)

    if "upload" in raw_path and http_method == "POST":
        return _handle_upload(event, request_id, timestamp)

    if "chat" in raw_path and http_method == "POST":
        return _handle_chat(event, request_id, timestamp)

    return error(404, "Endpoint no encontrado", request_id)


def _handle_upload(event, request_id: str, timestamp: str) -> dict:
    try:
        body = json.loads(event.get("body", "{}"))
        juzgado = body.get("juzgado", "")

        files_data = body.get("files", [])
        if not files_data:
            single_file = body.get("file", "")
            single_name = body.get("filename", "document.pdf")
            if single_file:
                files_data = [{"file": single_file, "filename": single_name}]

        if not files_data:
            return error(400, "Se requiere al menos un archivo", request_id)

        import base64

        extracted_files = []
        processed_files = []
        skipped_files = []

        for file_item in files_data:
            file_base64 = file_item.get("file", "")
            filename = file_item.get("filename", "document.pdf")

            if not file_base64:
                skipped_files.append({"filename": filename, "reason": "Archivo vacío"})
                continue

            ext = os.path.splitext(filename.lower())[1]
            if ext and ext not in SUPPORTED_EXTENSIONS:
                skipped_files.append({
                    "filename": filename,
                    "reason": f"Formato {ext} no soportado",
                })
                continue

            try:
                file_bytes = base64.b64decode(file_base64)
                text = extract_text(file_bytes, filename)
            except Exception as e:
                skipped_files.append({"filename": filename, "reason": f"Error: {str(e)}"})
                continue

            if not text.strip():
                skipped_files.append({"filename": filename, "reason": "No se pudo extraer texto"})
                continue

            processed_files.append(filename)
            extracted_files.append({"filename": filename, "text": text})

        if not processed_files:
            return error(400, "No se pudo extraer texto de ningún archivo", request_id)

        repo = _get_radicado_repo()
        all_radicados = repo.get_all_radicados()

        matches = SearchService.search(
            files=extracted_files,
            radicados=all_radicados,
            juzgado_filter=juzgado or None,
        )

        response_text = SearchService.format_response(matches)

        return success({
            "response": response_text,
            "matches": [m.to_dict() for m in matches],
            "total_matches": len(matches),
            "files_processed": processed_files,
            "skipped_files": skipped_files,
            "request_id": request_id,
            "timestamp": timestamp,
        })

    except ValueError as e:
        return error(400, str(e), request_id)
    except Exception:
        logger.error("Upload error", exc_info=True)
        return error(500, "Error procesando el archivo", request_id)


def _handle_chat(event, request_id: str, timestamp: str) -> dict:
    try:
        body = json.loads(event.get("body", "{}"))
        message = body.get("message", "").strip()

        if not message:
            return error(400, "El campo 'message' es requerido", request_id)

        repo = _get_radicado_repo()

        search_results = repo.search_radicado(None, message)

        if search_results:
            lines = [f"Se encontraron **{len(search_results)}** resultado(s):\n"]
            for r in search_results:
                lines.append(
                    f"- Radicado **{r.radicado}** (No. {r.numero}) "
                    f"— Juzgado: **{r.juzgado}** — {r.relacion}"
                )
            response_text = "\n".join(lines)
        else:
            name_results = repo.search_by_name(None, message)
            if name_results:
                lines = [f"Se encontraron **{len(name_results)}** resultado(s) por nombre:\n"]
                for r in name_results:
                    lines.append(
                        f"- Radicado **{r.radicado}** (No. {r.numero}) "
                        f"— Juzgado: **{r.juzgado}** — {r.relacion}"
                    )
                response_text = "\n".join(lines)
            else:
                response_text = (
                    "No se encontraron resultados para tu búsqueda.\n\n"
                    "Puedes:\n"
                    "- Buscar por número de radicado (ej: `0421`, `2024-00065`)\n"
                    "- Buscar por nombre (ej: `Rodriguez`)\n"
                    "- Subir un documento PDF/DOCX para buscar coincidencias"
                )

        return success({
            "response": response_text,
            "request_id": request_id,
            "timestamp": timestamp,
        })

    except Exception:
        logger.error("Chat error", exc_info=True)
        return error(500, "Error interno del servidor", request_id)
